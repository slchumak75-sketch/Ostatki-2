# -*- coding: utf-8 -*-
"""DEV-ONLY сервер: тот же API и тот же фронтенд, что у Go-приложения.

Нужен для: (1) живого превью GUI без установленного Go; (2) разработки фронтенда;
(3) сверки Go/Python реализаций алгоритма. Продуктовая сборка — Go (Ostatki.exe).
Запуск: python3 tools/dev_server.py [порт]   (по умолчанию 8765, 0.0.0.0)
При старте автоматически загружает *.xlsx из корня репозитория, если есть.
"""
import glob
import io
import json
import os
import sys
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import quote

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast import compute, summarize, DEFAULT_PARAMS  # noqa: E402
import read_xlsx  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WEB = os.path.join(ROOT, "internal", "app", "web")
VERSION = "py-dev"

_lock = threading.Lock()
_ctx = {"rows": None, "months": None, "meta": None, "state": None, "raw": None, "name": ""}

MIME = {".html": "text/html; charset=utf-8", ".css": "text/css; charset=utf-8",
        ".js": "text/javascript; charset=utf-8"}


def recompute(params):
    st = summarize(compute(_ctx["rows"], _ctx["months"], params), _ctx["rows"])
    st["fileName"] = _ctx["name"]
    st["version"] = VERSION
    st["loadedAt"] = time.strftime("%Y-%m-%d %H:%M:%S")
    _ctx["state"] = st
    return st


def load_bytes(name, data, params):
    tmp = os.path.join("/tmp", "ostatki_upload.xlsx")
    with open(tmp, "wb") as f:
        f.write(data)
    rows, months, meta = read_xlsx.load(tmp)
    with _lock:
        _ctx.update({"rows": rows, "months": months, "meta": meta, "raw": data, "name": name})
        return recompute(params or DEFAULT_PARAMS)


def parse_multipart(body, content_type):
    """Минимальный разбор multipart/form-data для полей file и params."""
    if "boundary=" not in content_type:
        return {}
    boundary = content_type.split("boundary=")[1].strip().strip('"').encode()
    out = {}
    for part in body.split(b"--" + boundary):
        if b"\r\n\r\n" not in part:
            continue
        head, payload = part.split(b"\r\n\r\n", 1)
        payload = payload.rstrip(b"\r\n")
        if payload.endswith(b"--"):
            payload = payload[:-2].rstrip(b"\r\n")
        name = filename = None
        for token in head.decode("utf-8", "replace").split(";"):
            token = token.strip()
            if token.startswith("name="):
                name = token[5:].strip('"')
            if token.startswith("filename="):
                filename = token[9:].strip('"')
        if name:
            out[name] = (filename, payload)
    return out


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, fmt, *args):
        sys.stderr.write("%s %s\n" % (self.command, self.path))

    def _send(self, code, data, ctype, extra=None):
        if isinstance(data, str):
            data = data.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(data)

    def _json(self, code, obj):
        self._send(code, json.dumps(obj, ensure_ascii=False), "application/json; charset=utf-8")

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        return self.rfile.read(n) if n else b""

    def do_GET(self):
        path = self.path.split("?")[0]
        try:
            if path == "/api/health":
                return self._json(200, {"ok": True, "version": VERSION})
            if path == "/api/state":
                with _lock:
                    st = _ctx["state"]
                return self._json(200, st if st else {"empty": True, "version": VERSION})
            if path == "/api/export.xlsx":
                return self._export_xlsx()
            if path == "/api/purchase.csv":
                return self._export_csv()
            name = "index.html" if path in ("/", "") else os.path.basename(path)
            full = os.path.join(WEB, name)
            if os.path.isfile(full):
                with open(full, "rb") as f:
                    return self._send(200, f.read(), MIME.get(os.path.splitext(name)[1], "application/octet-stream"))
            return self._json(404, {"ok": False, "error": "не найдено"})
        except Exception as e:  # noqa: BLE001
            return self._json(500, {"ok": False, "error": str(e)})

    def do_POST(self):
        path = self.path.split("?")[0]
        try:
            if path == "/api/upload":
                form = parse_multipart(self._body(), self.headers.get("Content-Type", ""))
                if "file" not in form:
                    return self._json(400, {"ok": False, "error": "в запросе нет файла (поле file)"})
                filename, data = form["file"]
                params = dict(DEFAULT_PARAMS)
                if "params" in form:
                    try:
                        params.update(json.loads(form["params"][1].decode("utf-8")))
                    except Exception:  # noqa: BLE001
                        pass
                try:
                    st = load_bytes(filename or "upload.xlsx", data, params)
                except Exception as e:  # noqa: BLE001
                    return self._json(422, {"ok": False, "error": str(e)})
                return self._json(200, st)
            if path == "/api/recalc":
                body = json.loads(self._body().decode("utf-8") or "{}")
                with _lock:
                    if _ctx["rows"] is None:
                        return self._json(400, {"ok": False, "error": "сначала загрузите файл ведомости"})
                    params = dict(DEFAULT_PARAMS)
                    params.update(body.get("params") or {})
                    st = recompute(params)
                return self._json(200, st)
            if path == "/api/shutdown":
                self._json(200, {"ok": True})
                threading.Timer(0.3, lambda: os._exit(0)).start()
                return
            return self._json(404, {"ok": False, "error": "не найдено"})
        except Exception as e:  # noqa: BLE001
            return self._json(500, {"ok": False, "error": str(e)})

    def _export_xlsx(self):
        with _lock:
            st, raw, meta = _ctx["state"], _ctx["raw"], _ctx["meta"]
        if not st or raw is None:
            return self._json(400, {"ok": False, "error": "нет загруженных данных"})
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb.worksheets[0]
        pcol = meta["purchaseCol"] or 8
        hdr = ws.cell(meta["headerRow"], pcol)
        hdr.value = "К закупке\n(рекомендация)"
        hdr.font = Font(bold=True, color="FFFFFF")
        hdr.fill = PatternFill("solid", fgColor="2F5496")
        hdr.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        fill = PatternFill("solid", fgColor="FFF2CC")
        for r in st["rows"]:
            c = ws.cell(r["_excelRow"], pcol)
            c.value = int(r["recommend"]) if float(r["recommend"]).is_integer() else r["recommend"]
            if r["recommend"] > 0:
                c.fill = fill
                c.font = Font(bold=True)
        buf = io.BytesIO()
        wb.save(buf)
        name = quote("Ведомость_К_закупке.xlsx")
        return self._send(200, buf.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            {"Content-Disposition": f"attachment; filename*=UTF-8''{name}"})

    def _export_csv(self):
        with _lock:
            st = _ctx["state"]
        if not st:
            return self._json(400, {"ok": False, "error": "нет загруженных данных"})

        def num(v):
            return f"{v:.2f}".replace(".", ",") if v is not None else ""

        def cell(s):
            s = str(s or "")
            return '"' + s.replace('"', '""') + '"' if any(ch in s for ch in ';"\n\r') else s

        lines = ["Поставщик;Артикул;Номенклатура;Ед.;Остаток;Прогноз на период;К закупке;Цена поставщика;Сумма заказа;Склад"]
        rows = sorted([r for r in st["rows"] if r["recommend"] > 0],
                      key=lambda r: (r["supplier"], -r["orderCost"]))
        for r in rows:
            lines.append(";".join([cell(r["supplier"]), cell(r["article"]), cell(r["name"]), cell(r["unit"]),
                                   num(r["stock"]), num(r["forecastTotal"]), num(r["recommend"]),
                                   num(r["purchasePrice"]), num(r["orderCost"]), cell(r["warehouse"])]))
        data = ("\ufeff" + "\r\n".join(lines)).encode("utf-8")
        name = quote("Закупка_по_поставщикам.csv")
        return self._send(200, data, "text/csv; charset=utf-8",
                            {"Content-Disposition": f"attachment; filename*=UTF-8''{name}"})


def main():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 8765
    files = [f for f in glob.glob(os.path.join(ROOT, "*.xlsx")) if "закуп" not in f.lower()]
    if files:
        try:
            with open(files[0], "rb") as f:
                load_bytes(os.path.basename(files[0]), f.read(), DEFAULT_PARAMS)
            print(f"Предзагружено: {files[0]}", flush=True)
        except Exception as e:  # noqa: BLE001
            print(f"Предзагрузка не удалась: {e}", flush=True)
    srv = ThreadingHTTPServer(("0.0.0.0", port), Handler)
    print(f"Ostatki DEV (python) -> http://127.0.0.1:{port}/", flush=True)
    srv.serve_forever()


if __name__ == "__main__":
    main()
