# -*- coding: utf-8 -*-
"""Чтение ведомости 1С (.xlsx) для Python-инструментов. Зеркало internal/xlsxio (Go)."""
import openpyxl
from forecast import parse_num, parse_month_header


def _s(v):
    return "" if v is None else str(v).strip()


def load(path):
    """-> (rows, months, meta). Строка заголовков ищется по словам, не по номеру."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header_row = None
    col = {}
    months = []
    month_cols = []
    for r in range(1, min(20, ws.max_row) + 1):
        texts = [_s(ws.cell(r, c).value).lower() for c in range(1, ws.max_column + 1)]
        if any("номенклатура" == t for t in texts) and any(t == "склад" for t in texts):
            header_row = r
            for c, t in enumerate(texts, start=1):
                raw = ws.cell(r, c).value
                m = parse_month_header(raw)
                if m:
                    months.append(m)
                    month_cols.append(c)
                elif t == "склад":
                    col["warehouse"] = c
                elif "группа" in t:
                    col["group"] = c
                elif t == "артикул":
                    col["article"] = c
                elif t == "номенклатура":
                    col["name"] = c
                elif t.startswith("ед"):
                    col["unit"] = c
                elif t == "расход":
                    col["rashod"] = c
                elif "остаток" in t:
                    col["stock"] = c
                elif "цена" in t and "постав" in t:
                    col["purchasePrice"] = c
                elif "рознич" in t:
                    col["retailPrice"] = c
                elif "поставщик" in t:
                    col["supplier"] = c
            break
    if header_row is None:
        raise ValueError("Не найдена строка заголовков (нужны «Склад» и «Номенклатура»)")
    # Столбец «К закупке»: пустой столбец между «Конечный остаток» и «Цена поставщика»
    purchase_col = None
    if "stock" in col and "purchasePrice" in col and col["purchasePrice"] == col["stock"] + 2:
        purchase_col = col["stock"] + 1

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = _s(ws.cell(r, col.get("name", 4)).value)
        if not name:
            continue
        sales = []
        for c in month_cols:
            v = parse_num(ws.cell(r, c).value)
            sales.append(v or 0.0)
        rows.append({
            "warehouse": _s(ws.cell(r, col["warehouse"]).value) if "warehouse" in col else "",
            "group": _s(ws.cell(r, col["group"]).value) if "group" in col else "",
            "article": _s(ws.cell(r, col["article"]).value) if "article" in col else "",
            "name": name,
            "unit": _s(ws.cell(r, col["unit"]).value) if "unit" in col else "",
            "rashod": parse_num(ws.cell(r, col["rashod"]).value) if "rashod" in col else None,
            "stock": parse_num(ws.cell(r, col["stock"]).value) if "stock" in col else None,
            "purchasePrice": parse_num(ws.cell(r, col["purchasePrice"]).value) if "purchasePrice" in col else None,
            "retailPrice": parse_num(ws.cell(r, col["retailPrice"]).value) if "retailPrice" in col else None,
            "supplier": _s(ws.cell(r, col["supplier"]).value) if "supplier" in col else "",
            "sales": sales,
            "_excelRow": r,
        })
    meta = {"headerRow": header_row, "columns": col, "monthCols": month_cols,
            "purchaseCol": purchase_col, "sheet": ws.title}
    return rows, months, meta
