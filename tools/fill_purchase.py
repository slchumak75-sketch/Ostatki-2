# -*- coding: utf-8 -*-
"""Заполняет столбец H (8-й, «К закупке») в КОПИИ исходной ведомости.

Оригинал не трогается. Дополнительно добавляются листы:
  «Закупка по поставщикам», «Сезонность», «Методика».
Запуск: python3 tools/fill_purchase.py [вход.xlsx] [выход.xlsx]
Параметры алгоритма — DEFAULT_PARAMS в tools/forecast.py (те же, что в Go-приложении).
"""
import glob
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from forecast import compute, summarize, DEFAULT_PARAMS, MONTH_NAMES_RU  # noqa: E402
from read_xlsx import load  # noqa: E402


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    src = sys.argv[1] if len(sys.argv) > 1 else \
        [f for f in glob.glob(os.path.join(root, "*.xlsx")) if "закуп" not in f.lower()][0]
    out_dir = os.path.join(root, "output")
    os.makedirs(out_dir, exist_ok=True)
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(out_dir, "Ведомость_К_закупке.xlsx")

    rows, months, meta = load(src)
    state = summarize(compute(rows, months, DEFAULT_PARAMS), rows)
    pcol = meta["purchaseCol"] or 8
    letter = get_column_letter(pcol)

    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    hdr = ws.cell(meta["headerRow"], pcol)
    hdr.value = "К закупке\n(рекомендация)"
    hdr.font = Font(bold=True, color="FFFFFF")
    hdr.fill = PatternFill("solid", fgColor="2F5496")
    hdr.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[letter].width = 13
    by_excel_row = {r["_excelRow"]: r for r in state["rows"]}
    fill_order = PatternFill("solid", fgColor="FFF2CC")
    for erow, r in by_excel_row.items():
        c = ws.cell(erow, pcol)
        c.value = int(r["recommend"]) if float(r["recommend"]).is_integer() else r["recommend"]
        if r["recommend"] > 0:
            c.fill = fill_order
            c.font = Font(bold=True)

    # --- лист «Закупка по поставщикам» ---
    ws2 = wb.create_sheet("Закупка по поставщикам")
    ws2.append(["Поставщик", "Артикул", "Номенклатура", "Ед.", "Остаток",
                "Прогноз на период", "К закупке", "Цена поставщика", "Сумма заказа", "Склад"])
    order_rows = sorted([r for r in state["rows"] if r["recommend"] > 0],
                          key=lambda x: (x["supplier"], -x["orderCost"]))
    for r in order_rows:
        ws2.append([r["supplier"], r["article"], r["name"], r["unit"], r["stock"],
                    r["forecastTotal"], r["recommend"], r["purchasePrice"],
                    r["orderCost"], r["warehouse"]])
    s = state["summary"]
    ws2.append([])
    ws2.append(["ИТОГО", "", f"{s['orderSku']} SKU", "", "", "",
                s["orderUnits"], "", s["orderBudget"], ""])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="548235")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:J{len(order_rows) + 1}"
    for col_letter, w in zip("ABCDEFGHIJ", [28, 12, 52, 6, 9, 11, 10, 11, 12, 16]):
        ws2.column_dimensions[col_letter].width = w

    # --- лист «Сезонность» ---
    ws3 = wb.create_sheet("Сезонность")
    ws3.append(["Месяц", "Все товары"] + list(state["seasonalByGroup"].keys()))
    for i in range(12):
        ws3.append([MONTH_NAMES_RU[i], state["seasonalGlobal"][i]] +
                   [state["seasonalByGroup"][g][i] for g in state["seasonalByGroup"]])
    for c in ws3[1]:
        c.font = Font(bold=True)
    ws3.column_dimensions["A"].width = 14

    # --- лист «Методика» ---
    ws4 = wb.create_sheet("Методика")
    fc_labels = ", ".join(m["label"] for m in state["forecastMonths"])
    lines = [
        ["Параметры и методика расчёта столбца «К закупке»"],
        [""],
        ["Период прогноза", fc_labels],
        ["Горизонт (полных месяцев)", DEFAULT_PARAMS["horizon"]],
        ["Страховой запас, %", DEFAULT_PARAMS["safetyPercent"]],
        ["Срок поставки, дней", DEFAULT_PARAMS["leadDays"]],
        ["Нет продаж N мес. -> не закупать", DEFAULT_PARAMS["deadMonths"]],
        ["Последний месяц неполный", str(DEFAULT_PARAMS["partialLast"])],
        [""],
        ["Формула: К закупке = max(0, ПрогнозСпроса + СпросЗаСрокПоставки + СтраховойЗапас − Остаток)"],
        ["Прогноз месяца = БазовыйУровень × ИндексСезонности(месяц)."],
        ["БазовыйУровень — среднее десезонализированное потребление за последние 12 полных месяцев (+ затухающий тренд)."],
        ["Индекс сезонности товара сглаживается к индексу его группы: чем меньше история товара, тем больше вес группы."],
        ["Полная методика: docs/ALGORITHM.md в репозитории. Тот же расчёт выполняет приложение Ostatki.exe."],
        [""],
        ["ИТОГО к закупке: SKU", s["orderSku"]],
        ["ИТОГО к закупке: единиц", s["orderUnits"]],
        ["ИТОГО бюджет закупки (по ценам поставщика)", s["orderBudget"]],
        [""],
        ["Предупреждения о данных:"],
    ] + [[" - " + w] for w in state["warnings"]]
    for line in lines:
        ws4.append(line)
    ws4["A1"].font = Font(bold=True, size=12)
    ws4.column_dimensions["A"].width = 60
    ws4.column_dimensions["B"].width = 70

    wb.save(dst)
    print(f"Сохранено: {dst}")
    print(f"К закупке: {s['orderSku']} SKU, {s['orderUnits']:.0f} ед., бюджет {s['orderBudget']:,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
